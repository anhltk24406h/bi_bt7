"""Chuyển bản Markdown báo cáo hiện tại thành Word có văn bản/bảng chỉnh sửa được.

Chạy: python export_docx.py. Sáu biểu đồ được nhúng dưới dạng ảnh PNG.
"""
from pathlib import Path
import re
from zipfile import ZipFile
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent


def export():
    source = ROOT / 'outputs/bao_cao_visual_analytics.md'
    lines = source.read_text(encoding='utf-8').splitlines()
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.6), Cm(1.6)
    section.left_margin, section.right_margin = Cm(1.6), Cm(1.6)
    for style_name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
        style = doc.styles[style_name]
        style.font.name = 'Arial'
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Arial')
    normal = doc.styles['Normal']
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string('253B4C')
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.widow_control = True
    for name, size in [('Title', 24), ('Subtitle', 15), ('Heading 1', 16)]:
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].font.color.rgb = RGBColor.from_string('176B87')
    doc.core_properties.title = 'Doanh thu tập trung ở đâu? — UCI Online Retail'
    doc.core_properties.subject = 'Visual Analytics và Data Storytelling'
    doc.core_properties.author = 'Nhóm sinh viên'
    footer = section.footer.paragraphs[0]
    footer.text = 'UCI ONLINE RETAIL  |  VISUAL ANALYTICS                                      '
    footer.style = 'Normal'
    for run in footer.runs:
        run.font.size = Pt(8)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)

    i = 0
    while i < len(lines):
        text = lines[i].strip()
        if not text:
            i += 1
            continue
        if text == '---':
            doc.add_page_break()
        elif text.startswith('## '):
            doc.add_heading(text[3:], level=1)
        elif text.startswith('|'):
            header = text.strip('|').split('|')
            cols = len(header)
            table = doc.add_table(rows=0, cols=cols)
            table.style = 'Light Shading Accent 1'
            while i < len(lines) and lines[i].strip().startswith('|'):
                row_text = lines[i].strip().strip('|')
                if not re.fullmatch(r'[\s|:\-]+', row_text):
                    values = row_text.split('|', maxsplit=cols - 1)
                    cells = table.add_row().cells
                    for cell, value in zip(cells, values):
                        cell.text = value.strip()
                        for paragraph in cell.paragraphs:
                            paragraph.paragraph_format.space_after = Pt(5)
                            for run in paragraph.runs:
                                run.font.size = Pt(9)
                    no_split = OxmlElement('w:cantSplit')
                    table.rows[-1]._tr.get_or_add_trPr().append(no_split)
                i += 1
            repeat = OxmlElement('w:tblHeader')
            table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for cell in table.rows[0].cells:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue
        elif text.startswith('!['):
            match = re.fullmatch(r'!\[(.*?)\]\((.*?)\)', text)
            assert match, text
            image_path = source.parent / match.group(2)
            doc.add_picture(str(image_path), width=Cm(17.7))
            p = doc.paragraphs[-1]
            p.paragraph_format.space_after = Pt(8)
            # Alternative text preserved for screen readers.
            drawing = p._p.xpath('.//wp:docPr')
            if drawing:
                drawing[0].set('descr', match.group(1))
        elif text == 'Doanh thu tập trung ở đâu?':
            doc.add_paragraph(text, style='Title')
        elif text == 'Khách hàng, sản phẩm và giao dịch hủy trong UCI Online Retail':
            doc.add_paragraph(text, style='Subtitle')
        else:
            p = doc.add_paragraph()
            if text.startswith(('Phát hiện — ', 'Diễn giải — ', 'Hành động — ')):
                label, body = text.split(' — ', 1)
                p.add_run(label + ' — ').bold = True
                p.add_run(body)
            else:
                p.add_run(text)
        i += 1

    target = ROOT / 'outputs/bao_cao_visual_analytics.docx'
    doc.save(target)
    reopened = Document(target)
    assert len(reopened.inline_shapes) == 6
    assert len(reopened.tables) == 3
    text = '\n'.join(p.text for p in reopened.paragraphs)
    assert 'Doanh thu tập trung ở đâu?' in text
    assert '10,666,684.54' in text and '27.39%' in text
    with ZipFile(target) as z:
        assert z.testzip() is None
        assert len([n for n in z.namelist() if n.startswith('word/media/')]) == 6
        assert z.read('word/document.xml').count(b'w:type="page"') == 9
    print('DOCX verified: 6 images, 3 editable tables, 9 page breaks.')
    print(target)


if __name__ == '__main__':
    export()
